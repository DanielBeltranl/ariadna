"""Shared sheet writer for the Excel generators: one sheet per module, styled header, wrapped rows."""
import math
import re

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

FIRST_ROW, FIRST_COL = 2, 2  # leave column A and row 1 as a margin, like the reference workbooks
BORDER = Border(*(Side(style="thin", color="D9D9D9"),) * 4)
LINE_HEIGHT = 15


def sheet_title(name, used):
    """Excel sheet names: max 31 chars, no []:*?/\\ and unique (case-insensitive)."""
    base = re.sub(r"[\[\]:*?/\\]", " ", name or "General").strip()[:31] or "General"
    title, n = base, 2
    while title.lower() in used:
        suffix = f" ({n})"
        title, n = base[:31 - len(suffix)] + suffix, n + 1
    used.add(title.lower())
    return title


def group_by_module(items):
    """Keep modules in first-appearance order."""
    groups = {}
    for item in items:
        groups.setdefault(item.get("modulo") or "General", []).append(item)
    return groups


def numbered(lines):
    return "\n\n".join(f"{i}. {line}" for i, line in enumerate(lines, 1))


def _row_height(values, widths):
    lines = 1
    for value, width in zip(values, widths):
        text_lines = str(value).split("\n")
        lines = max(lines, sum(max(1, math.ceil(len(t) / (width * 1.1))) for t in text_lines))
    return min(lines * LINE_HEIGHT, 409)


def write_table(ws, headers, rows, widths, header_fill, header_font_color, stripe_fills=None):
    header_font = Font(name="Calibri", size=11, bold=True, color=header_font_color)
    body_font = Font(name="Calibri", size=10, color="1F1F1F")
    for offset, width in enumerate(widths):
        ws.column_dimensions[ws.cell(row=1, column=FIRST_COL + offset).column_letter].width = width
    ws.column_dimensions["A"].width = 3

    for offset, header in enumerate(headers):
        cell = ws.cell(row=FIRST_ROW, column=FIRST_COL + offset, value=header)
        cell.font, cell.border = header_font, BORDER
        cell.fill = PatternFill("solid", fgColor=header_fill)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[FIRST_ROW].height = 22

    for r, values in enumerate(rows, FIRST_ROW + 1):
        fill = PatternFill("solid", fgColor=stripe_fills[(r - FIRST_ROW - 1) % 2]) if stripe_fills else None
        for offset, value in enumerate(values):
            cell = ws.cell(row=r, column=FIRST_COL + offset, value=value)
            cell.font, cell.border = body_font, BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if fill:
                cell.fill = fill
        ws.row_dimensions[r].height = _row_height(values, widths)
    ws.freeze_panes = ws.cell(row=FIRST_ROW + 1, column=FIRST_COL)
