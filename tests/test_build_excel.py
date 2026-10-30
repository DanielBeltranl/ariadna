import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from xlsx_sheets import sheet_title  # noqa: E402


def run(script, spec, tmp):
    spec_path, out = Path(tmp) / "spec.json", Path(tmp) / "out.xlsx"
    spec_path.write_text(json.dumps(spec), encoding="utf-8")
    proc = subprocess.run([sys.executable, str(ROOT / "scripts" / script), "--json", str(spec_path), "--output", str(out)],
                          capture_output=True, text=True, encoding="utf-8", env={**os.environ, "PYTHONIOENCODING": "utf-8"})
    return proc, out


def example(name):
    return json.loads((ROOT / "examples" / name).read_text(encoding="utf-8"))


class CasosUsoTest(unittest.TestCase):
    def test_one_sheet_per_module_with_reference_columns(self):
        with tempfile.TemporaryDirectory() as tmp:
            proc, out = run("build_casos_uso.py", example("casos.json"), tmp)
            self.assertEqual(proc.returncode, 0, proc.stderr)
            wb = load_workbook(out)
            self.assertEqual(wb.sheetnames, ["Detección", "Registro de tickets"])
            ws = wb["Detección"]
            self.assertEqual([c.value for c in ws[2][1:7]],
                             ["ID", "Caso de Uso", "Actor", "Precondiciones", "Flujo de Pasos", "Resultado Esperado"])
            self.assertEqual(ws["B3"].value, "UC-01")
            self.assertTrue(ws["F3"].value.startswith("1. El sistema consulta"))
            self.assertIn("\n\n2. ", ws["F3"].value)
            self.assertEqual(ws["B2"].fill.fgColor.rgb, "00535FC1")

    def test_rejects_case_without_required_field(self):
        spec = example("casos.json")
        del spec["casos"][0]["actor"]
        with tempfile.TemporaryDirectory() as tmp:
            proc, out = run("build_casos_uso.py", spec, tmp)
            self.assertFalse(out.exists())
        self.assertEqual(proc.returncode, 1)
        self.assertIn("actor", proc.stderr)


class HistoriasTest(unittest.TestCase):
    def test_story_format_matches_reference(self):
        with tempfile.TemporaryDirectory() as tmp:
            proc, out = run("build_historias.py", example("historias.json"), tmp)
            self.assertEqual(proc.returncode, 0, proc.stderr)
            ws = load_workbook(out)["Detección"]
            self.assertEqual([c.value for c in ws[2][1:5]], ["N°", "Nombre", "Descripción", "Criterios de Aceptación"])
            self.assertEqual(ws["B3"].value, "1.0")
            self.assertEqual(ws["D3"].value, "Como operador N1, quiero que las caídas de enlace se detecten "
                                             "automáticamente para no revisar Zabbix manualmente.")
            self.assertTrue(ws["E3"].value.startswith("1. El sistema debe consultar"))
            self.assertEqual(ws["B2"].fill.fgColor.rgb, "00A4C2F4")

    def test_rejects_empty_list(self):
        with tempfile.TemporaryDirectory() as tmp:
            proc, _ = run("build_historias.py", {"historias": []}, tmp)
        self.assertEqual(proc.returncode, 1)


class SheetTitleTest(unittest.TestCase):
    def test_sanitizes_truncates_and_deduplicates(self):
        used = set()
        self.assertEqual(sheet_title("Alertas/Monitoreo", used), "Alertas Monitoreo")
        long_name = "x" * 40
        self.assertEqual(len(sheet_title(long_name, used)), 31)
        self.assertEqual(sheet_title(long_name, used), "x" * 27 + " (2)")
        self.assertEqual(sheet_title("", used), "General")


if __name__ == "__main__":
    unittest.main()
